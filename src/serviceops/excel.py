"""Excel report: the reporting views as formatted tables with native Excel charts."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import queries as q
from .analysis import findings
from .db import Database
from .taxonomy import SNAPSHOT

INK = "16181A"
ACCENT = "1A6F57"
MUTED = "686C71"
PCT = '0.0"%"'


def _table(ws, name: str, rows: list[dict], columns: list[tuple[str, str, str | None]],
           top: int = 4, left: int = 1) -> tuple[int, int]:
    """Write rows as an Excel table. columns: (key, header, number_format)."""
    for j, (_, header, _) in enumerate(columns):
        ws.cell(row=top, column=left + j, value=header)
    for i, row in enumerate(rows, 1):
        for j, (key, _, fmt) in enumerate(columns):
            cell = ws.cell(row=top + i, column=left + j, value=row.get(key))
            if fmt:
                cell.number_format = fmt
    last = top + max(len(rows), 1)
    ref = f"{get_column_letter(left)}{top}:{get_column_letter(left + len(columns) - 1)}{last}"
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    for j, (key, header, _) in enumerate(columns):
        width = max([len(header)] + [len(str(r.get(key, ""))) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(left + j)].width = min(max(width, 10), 48)
    return top, last


def _title(ws, title: str, source: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(size=15, bold=True, color=INK)
    ws["A2"] = source
    ws["A2"].font = Font(size=9, italic=True, color=MUTED)
    ws.sheet_view.showGridLines = False


def _chart(chart, title: str, y_title: str | None = None, width: float = 18, height: float = 8):
    chart.title = title
    chart.style = 10
    chart.width, chart.height = width, height
    chart.legend = None
    if y_title:
        chart.y_axis.title = y_title
    return chart


def build_workbook(db: Database, path: Path, quality_report: dict) -> Path:
    f = q.Filters()
    wb = Workbook()
    fnd = findings(db, f)
    k = fnd["kpis"]

    # Summary ---------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    _title(ws, "IT Service Ops & SLA Report",
           f"Simulated service desk, tickets opened Jan 2024 to Jun 2026. Snapshot {SNAPSHOT}. "
           "Source: vw_kpi_summary, vw_category_pareto.")
    cards = [("Tickets analysed", k["total_tickets"], "#,##0"),
             ("SLA compliance", k["sla_compliance_pct"], PCT),
             ("First-response compliance", k["response_compliance_pct"], PCT),
             ("SLA breaches", k["sla_breaches"], "#,##0"),
             ("Median resolution (h)", k["median_resolution_hours"], "0.0"),
             ("90th pct resolution (h)", k["p90_resolution_hours"], "0.0"),
             ("Open backlog", k["open_backlog"], "#,##0"),
             ("Reassigned at least once", k["reassigned_pct"], PCT)]
    for i, (label, value, fmt) in enumerate(cards):
        r, c = 4 + (i // 4) * 3, 1 + (i % 4) * 2
        ws.cell(row=r, column=c, value=label).font = Font(size=9, color=MUTED)
        cell = ws.cell(row=r + 1, column=c, value=value)
        cell.font = Font(size=18, bold=True, color=INK)
        cell.number_format = fmt
    for c in range(1, 9):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws["A11"] = "Finding"
    ws["A11"].font = Font(bold=True, color=ACCENT)
    ws["A12"] = fnd["headline"]
    ws["A12"].font = Font(size=12, bold=True)
    row = 14
    ws.cell(row=row, column=1, value="Recommendations").font = Font(bold=True, color=ACCENT)
    for i, rec in enumerate(fnd["recommendations"], 1):
        row += 1
        ws.cell(row=row, column=1, value=f"{i}. {rec['title']}").font = Font(bold=True)
        for label in ("evidence", "action", "owner"):
            row += 1
            ws.cell(row=row, column=1, value=f"{label.title()}: {rec[label]}")
            ws.cell(row=row, column=1).alignment = Alignment(wrap_text=False)
        row += 1

    # Monthly ---------------------------------------------------------------------------------
    ws = wb.create_sheet("Monthly")
    _title(ws, "Monthly volume and SLA compliance", "Source: vw_monthly_kpis (by opened month)")
    rows = q.monthly(db, f)
    top, last = _table(ws, "Monthly", rows, [
        ("opened_month", "Month", None), ("tickets_opened", "Tickets", "#,##0"),
        ("sla_breaches", "Breaches", "#,##0"), ("sla_compliance_pct", "SLA compliance %", PCT),
        ("response_compliance_pct", "Response compliance %", PCT),
        ("avg_resolution_hours", "Avg resolution (h)", "0.0")])
    cats = Reference(ws, min_col=1, min_row=top + 1, max_row=last)
    bar = _chart(BarChart(), "Tickets opened per month")
    bar.add_data(Reference(ws, min_col=2, min_row=top, max_row=last), titles_from_data=True)
    bar.set_categories(cats)
    ws.add_chart(bar, "H4")
    line = _chart(LineChart(), "SLA compliance by month", "%")
    line.add_data(Reference(ws, min_col=4, min_row=top, max_row=last), titles_from_data=True)
    line.set_categories(cats)
    ws.add_chart(line, "H22")

    # Pareto ----------------------------------------------------------------------------------
    ws = wb.create_sheet("Breach Pareto")
    _title(ws, "SLA breaches by category", "Source: vw_category_pareto")
    rows = q.pareto(db, f)
    top, last = _table(ws, "Pareto", rows, [
        ("breach_rank", "Rank", None), ("label", "Category", None),
        ("tickets", "Tickets", "#,##0"), ("breaches", "Breaches", "#,##0"),
        ("breach_rate_pct", "Breach rate %", PCT),
        ("share_of_volume_pct", "Share of volume %", PCT),
        ("share_of_breaches_pct", "Share of breaches %", PCT),
        ("cumulative_share_pct", "Cumulative share %", PCT)])
    bar = _chart(BarChart(), "Breaches by category")
    bar.type = "bar"
    bar.add_data(Reference(ws, min_col=4, min_row=top, max_row=last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=top + 1, max_row=last))
    bar.y_axis.scaling.orientation = "minMax"
    bar.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(bar, "J4")
    for r in range(top + 1, top + 3):
        for c in range(1, 9):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="E6F0EC")

    # Priority --------------------------------------------------------------------------------
    ws = wb.create_sheet("Priority SLA")
    _title(ws, "Compliance against each priority's target", "Source: vw_priority_sla")
    _table(ws, "PrioritySLA", q.by_priority(db, f), [
        ("priority_code", "Code", None), ("priority", "Priority", None),
        ("response_target_hours", "Response target (h)", "0.00"),
        ("resolution_target_hours", "Resolution target (h)", "0"),
        ("tickets", "Tickets", "#,##0"), ("breaches", "Breaches", "#,##0"),
        ("sla_compliance_pct", "SLA compliance %", PCT),
        ("response_compliance_pct", "Response compliance %", PCT),
        ("avg_resolution_hours", "Avg resolution (h)", "0.0")])

    # Drivers ---------------------------------------------------------------------------------
    ws = wb.create_sheet("Breach Drivers")
    _title(ws, "Third-party waits and hand-offs in the two highest-breach categories",
           "Source: vw_breach_drivers")
    _table(ws, "Drivers", q.drivers(db, f, fnd["top_categories"]), [
        ("category", "Category", None), ("waited_on_third_party", "Waited on third party", None),
        ("hop_bucket", "Hand-offs", None), ("tickets", "Tickets", "#,##0"),
        ("breaches", "Breaches", "#,##0"), ("breach_rate_pct", "Breach rate %", PCT)])

    # Backlog ---------------------------------------------------------------------------------
    ws = wb.create_sheet("Backlog")
    _title(ws, "Daily flow and open backlog", "Source: vw_daily_backlog, vw_backlog_aging")
    rows = q.backlog(db, f)
    top, last = _table(ws, "Backlog", rows, [
        ("flow_date", "Date", None), ("opened", "Opened", "#,##0"),
        ("resolved", "Resolved", "#,##0"), ("backlog", "Open at end of day", "#,##0")])
    line = _chart(LineChart(), "Open backlog at end of day", width=24)
    line.add_data(Reference(ws, min_col=4, min_row=top, max_row=last), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
    line.x_axis.number_format = "mmm yy"
    ws.add_chart(line, "G4")
    _table(ws, "Aging", q.aging(db, f), [
        ("age_band", "Age of open ticket", None), ("open_tickets", "Open", "#,##0"),
        ("already_breached", "Already breached", "#,##0")], top=4, left=20)

    # Groups & agents -------------------------------------------------------------------------
    ws = wb.create_sheet("Groups")
    _title(ws, "Resolver group workload", "Source: vw_group_workload")
    _table(ws, "Groups", q.groups(db, f), [
        ("assignment_group", "Group", None), ("tickets", "Tickets", "#,##0"),
        ("open_tickets", "Open", "#,##0"), ("breaches", "Breaches", "#,##0"),
        ("sla_compliance_pct", "SLA compliance %", PCT),
        ("avg_resolution_hours", "Avg resolution (h)", "0.0"),
        ("avg_reassignments", "Avg hand-offs", "0.00")])
    ws = wb.create_sheet("Agents")
    _title(ws, "Agent scorecard", "Source: vw_agent_scorecard")
    _table(ws, "Agents", q.agents(db, f), [
        ("assigned_agent", "Agent", None), ("assignment_group", "Group", None),
        ("tickets", "Tickets", "#,##0"), ("open_tickets", "Open", "#,##0"),
        ("breaches", "Breaches", "#,##0"), ("sla_compliance_pct", "SLA compliance %", PCT),
        ("avg_resolution_hours", "Avg resolution (h)", "0.0")])

    # Data quality ----------------------------------------------------------------------------
    ws = wb.create_sheet("Data Quality")
    _title(ws, "Data-quality gate",
           f"{quality_report['rows_in']:,} exported rows in, {quality_report['rows_clean']:,} "
           f"clean tickets out, {quality_report['rows_quarantined']:,} quarantined. "
           "Source: dq_check_results")
    _table(ws, "Checks", quality_report["checks"], [
        ("id", "Check", None), ("dimension", "Dimension", None), ("name", "Name", None),
        ("rule", "Rule", None), ("action", "Action", None), ("rows", "Rows", "#,##0")])
    quarantine = db.query("SELECT ticket_id, dq_reason FROM dq_quarantine ORDER BY dq_reason, "
                          "ticket_id")
    _table(ws, "Quarantine", quarantine, [("ticket_id", "Quarantined ticket", None),
                                          ("dq_reason", "Reason", None)], left=9)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
