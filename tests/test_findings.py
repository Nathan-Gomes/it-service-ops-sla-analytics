import pytest
from openpyxl import load_workbook


def test_headline_finding(built):
    f = built.findings
    assert f["top_categories"] == ["Access & Identity", "Network & VPN"]
    assert f["top_share_of_breaches_pct"] == pytest.approx(61, abs=0.5)
    assert f["top_share_of_volume_pct"] < 40
    assert "61% of SLA breaches" in f["headline"]


def test_drivers_support_the_recommendations(built):
    for item in built.findings["focus"]:
        assert item["breach_rate_waited_pct"] > 2 * item["breach_rate_no_wait_pct"]
        assert item["breach_rate_2plus_hops_pct"] > item["breach_rate_0_hops_pct"]
    assert len(built.findings["recommendations"]) == 3


def test_outputs_written(built):
    out = built.out_dir
    for name in ["serviceops.db", "tickets_raw_export.csv", "quarantine.csv",
                 "quality_report.json", "findings.json", "service_ops_sla_report.xlsx",
                 "powerbi/fact_ticket.csv", "powerbi/dim_date.csv", "powerbi/dim_priority.csv",
                 "powerbi/dim_category.csv", "powerbi/dq_checks.csv"]:
        assert (out / name).exists(), name


def test_excel_workbook(built):
    wb = load_workbook(built.out_dir / "service_ops_sla_report.xlsx")
    assert wb.sheetnames == ["Summary", "Monthly", "Breach Pareto", "Priority SLA",
                             "Breach Drivers", "Backlog", "Groups", "Agents", "Data Quality"]
    assert len(wb["Monthly"]._charts) == 2
    assert wb["Summary"]["A12"].value == built.findings["headline"]
    pareto = wb["Breach Pareto"]
    assert pareto["B5"].value == "Access & Identity"
    assert "Pareto" in pareto.tables
